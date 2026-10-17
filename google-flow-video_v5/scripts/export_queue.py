#!/usr/bin/env python3
"""Export one Excel sheet into the files used by queue mode.

usage: python export_queue.py <xlsx> <sheet> <out_dir> [--start N] [--no-header-split]

Creates in <out_dir> (one folder per sheet, e.g. Dev/):
  scenes.json   every row: row, scenario (forward-filled), scene, name, script, image   (used by mover_daemon.ps1)
  queue.json    [[name, image, tail], ...] from row index --start (default 0)  (what you inject with loadQueue)
  logs/ videos/ created empty
Prints: item count, header length, and the checksum [count, chars, sum-of-char-codes] of name+image+tail.
After injecting, window.queueChecksum() must print the SAME three numbers.

The common header: when every script starts with the same long block (the rule block already written in the Script
column), it is detected as the longest common prefix, cut at a paragraph end, stored in header.json and removed from
each tail, so only the short tails have to be pasted into the page. Use --no-header-split to keep full scripts.
Columns are found by header name: Scenario, Scene, Name of video, Script, Selected image.
Column A (Scenario) may be merged cells -> forward-filled.
"""
import json, os, sys

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl")


def main():
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    xlsx, sheet, out = a[0], a[1], a[2]
    start = int(sys.argv[sys.argv.index("--start") + 1]) if "--start" in sys.argv else 0
    split = "--no-header-split" not in sys.argv

    wb = openpyxl.load_workbook(xlsx)
    ws = wb[sheet]
    head = {str(c.value).strip().lower(): i + 1 for i, c in enumerate(ws[1]) if c.value}
    col = lambda *names: next((head[n] for n in names if n in head), None)
    c_sc, c_scene, c_name, c_script, c_img = (col("scenario"), col("scene"), col("name of video"), col("script"), col("selected image"))
    if not (c_script and c_name):
        sys.exit("need columns 'Name of video' and 'Script'; found: %s" % list(head))

    rows, last = [], None
    for r in range(2, ws.max_row + 1):
        script = ws.cell(r, c_script).value
        sc = (ws.cell(r, c_sc).value if c_sc else None) or last
        last = sc or last
        if not script:
            continue
        rows.append(dict(row=r, scenario=str(sc).strip() if sc else "", scene=ws.cell(r, c_scene).value if c_scene else "",
                         name=str(ws.cell(r, c_name).value).strip(),
                         image=(str(ws.cell(r, c_img).value).strip() if c_img and ws.cell(r, c_img).value else ""),
                         script=script))
    os.makedirs(os.path.join(out, "logs"), exist_ok=True)
    os.makedirs(os.path.join(out, "videos"), exist_ok=True)
    json.dump(rows, open(os.path.join(out, "scenes.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    names = [r["name"] for r in rows]
    dup = len(names) - len(set(names))
    header = ""
    if split and rows:
        p = os.path.commonprefix([r["script"] for r in rows])
        cut = p.rfind("\n\n")
        header = p[:cut] if cut > 200 else ""
    json.dump({"header": header}, open(os.path.join(out, "header.json"), "w", encoding="utf-8"), ensure_ascii=False)
    q = [[r["name"], r["image"], r["script"][len(header):]] for r in rows[start:]]
    json.dump(q, open(os.path.join(out, "queue.json"), "w", encoding="utf-8"), ensure_ascii=False)
    s = "".join(x[0] + x[1] + x[2] for x in q)
    print("items=%d header_chars=%d dup_names=%d" % (len(q), len(header), dup))
    print("CHECKSUM", [len(q), len(s), sum(ord(c) for c in s)])
    print("images:", sorted({r["image"] for r in rows}))
    print("scenarios:", sorted({r["scenario"] for r in rows}))


if __name__ == "__main__":
    main()

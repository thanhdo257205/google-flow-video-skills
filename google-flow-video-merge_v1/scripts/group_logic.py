#!/usr/bin/env python3
"""Luật ghép video theo tên (cột 'Name of video'). Chạy riêng để xem nhóm:
   python group_logic.py <file.xlsx> "<tên sheet>" [--scenario "Scenario 3"] [--json]
"""
import argparse, collections, json, re, sys

PAT = re.compile(r'^(\d+)([ABC]?)\.(\d+)([A-Z]?)$')   # <số><A/B/C?>.<thứ tự><chữ đuôi?>

def group(names):
    """names: tên video của MỘT scenario theo thứ tự trong sheet -> (groups, warnings).
    groups: list các dict {names:[...], key:(số đầu, chữ)}; mỗi nhóm = một video ghép."""
    warnings = []
    units = collections.OrderedDict()                  # (số, chữ, thứ tự) -> [(đuôi chữ, tên)]
    for n in names:
        m = PAT.match(n)
        if not m:
            raise ValueError('Tên video không đúng định dạng <số><A/B/C>.<thứ tự><đuôi chữ>: %r' % n)
        units.setdefault((int(m[1]), m[2], int(m[3])), []).append((m[4], n))   # luật 4: đuôi chữ = đoạn con
    groups = []
    for L in ['', 'A', 'B', 'C']:
        nums = sorted({k[0] for k in units if k[1] == L})
        comps = []
        for x in nums:
            if L and comps and x == comps[-1][-1] + 1: comps[-1].append(x)   # luật 3: hai số liền nhau cùng chữ
            else: comps.append([x])                                          # luật 1, 2: cùng số đầu
        for c in comps:
            if len(c) > 2:
                warnings.append('Chuỗi %d số liên tiếp cùng chữ %s: %s (luật chỉ nói "hai số", cần hỏi user)' % (len(c), L, c))
            keys = sorted([k for k in units if k[1] == L and k[0] in c], key=lambda k: (k[0], k[2]))
            ordered = [n for k in keys for _, n in sorted(units[k])]
            groups.append({'names': ordered, 'key': (c[0], L)})
    groups.sort(key=lambda g: g['key'])
    for g in groups:
        if len(g['names']) == 1:
            warnings.append('Video đứng một mình (không có gì để ghép): %s' % g['names'][0])
    # cảnh báo mẫu cần hỏi: đuôi chữ trong cùng nhóm với thứ tự khác (vd 34A.1A, 34A.1B, 34A.2)
    for k, v in units.items():
        if len(v) > 1 and k[1] != '' and len([1 for kk in units if kk[0] == k[0] and kk[1] == k[1]]) > 1:
            warnings.append('Đoạn con đuôi chữ lẫn với thứ tự khác trong cùng nhóm: %s (cần hỏi user)' % [n for _, n in v])
    return groups, warnings

def read_sheet(xlsx, sheet):
    """Trả OrderedDict scenario -> list (tên video, script). Cột tìm theo tiêu đề dòng 1; ô Scenario gộp được điền tiếp."""
    import openpyxl
    ws = openpyxl.load_workbook(xlsx, data_only=True)[sheet]
    head = {str(c.value).strip().lower(): i + 1 for i, c in enumerate(ws[1]) if c.value}
    need = {k: head.get(k) for k in ('scenario', 'name of video', 'script')}
    if None in need.values():
        raise ValueError('Sheet %r thiếu cột %s (tiêu đề hiện có: %s)' % (sheet, [k for k, v in need.items() if v is None], list(head)))
    out, last = collections.OrderedDict(), None
    for r in range(2, ws.max_row + 1):
        sc = ws.cell(r, need['scenario']).value or last; last = sc
        nm = ws.cell(r, need['name of video']).value
        if sc and nm:
            out.setdefault(str(sc).strip(), []).append((str(nm).strip(), str(ws.cell(r, need['script']).value or '')))
    return out

if __name__ == '__main__':
    ap = argparse.ArgumentParser(); ap.add_argument('xlsx'); ap.add_argument('sheet'); ap.add_argument('--scenario'); ap.add_argument('--json', action='store_true')
    a = ap.parse_args(); data = read_sheet(a.xlsx, a.sheet); res = {}
    for sc, items in data.items():
        if a.scenario and sc != a.scenario: continue
        gs, w = group([n for n, _ in items]); res[sc] = {'groups': [g['names'] for g in gs], 'warnings': w}
    if a.json: print(json.dumps(res, ensure_ascii=False, indent=1))
    else:
        for sc, r in res.items():
            print('==', sc)
            for g in r['groups']: print('  ', ' + '.join(g))
            for w in r['warnings']: print('   ! ' + w)
